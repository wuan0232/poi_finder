import json
from datetime import datetime
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import case, func, or_, select
from sqlalchemy.orm import Session

from database import get_db
from models import Poi, Task
from schemas import PoiPage, PoiProductUpdate, PoiRead, TaskCreate, TaskRead
from services.collect import collect_task
from services.cache import find_cached_task
from services.config import settings
from services.geo import wgs84_to_gcj02
from services.map_client import AmapClient, MapApiError, create_client


router = APIRouter(prefix="/api/tasks", tags=["tasks"])
pois_router = APIRouter(prefix="/api", tags=["pois"])


def _task_read(task: Task, cache_hit: bool = False) -> TaskRead:
    return TaskRead(
        id=task.id,
        center_name=task.center_name,
        center_address=task.center_address,
        keyword=task.keyword,
        types_code=task.types_code,
        radius_m=task.radius_m,
        center_lng=task.center_lng,
        center_lat=task.center_lat,
        city=task.city,
        keywords=json.loads(task.keywords_json),
        providers=json.loads(task.providers_json),
        status=task.status,
        total=task.total,
        error_message=task.error_message,
        cache_hit=cache_hit,
        created_at=task.created_at,
        updated_at=task.updated_at,
    )


def _poi_read(poi: Poi) -> PoiRead:
    return PoiRead.model_validate(poi)


@pois_router.get("/location/reverse")
async def reverse_location(
    longitude: float = Query(..., ge=-180, le=180),
    latitude: float = Query(..., ge=-90, le=90),
):
    if not settings.amap_key:
        raise HTTPException(status_code=400, detail="高德地图 Web 服务 Key 尚未配置")
    gcj_longitude, gcj_latitude = wgs84_to_gcj02(longitude, latitude)
    map_client = create_client("amap")
    if not isinstance(map_client, AmapClient):
        raise HTTPException(status_code=500, detail="高德地图客户端不可用")
    try:
        async with httpx.AsyncClient(timeout=settings.request_timeout) as client:
            result = await map_client.reverse_geocode(
                client, round(gcj_longitude, 6), round(gcj_latitude, 6)
            )
    except MapApiError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    return {
        **result,
        "longitude": round(gcj_longitude, 6),
        "latitude": round(gcj_latitude, 6),
    }


@router.post("", response_model=TaskRead, status_code=202)
def create_task(payload: TaskCreate, background_tasks: BackgroundTasks, db: Session = Depends(get_db)):
    has_coordinates = payload.center_lng is not None and payload.center_lat is not None
    is_around_search = bool(payload.center_name or payload.center_address or has_coordinates)
    if is_around_search and payload.providers != ["amap"]:
        raise HTTPException(
            status_code=400,
            detail="半径周边检索当前仅支持高德地图，请只选择 amap",
        )
    effective_providers = list(payload.providers)
    if (
        is_around_search
        and settings.website_enrichment_enabled
        and "official_website" not in effective_providers
    ):
        effective_providers.append("official_website")
    if is_around_search:
        effective_providers.append("healthcare_filter_v2")
    center_lng = center_lat = None
    if has_coordinates:
        center_lng, center_lat = wgs84_to_gcj02(payload.center_lng, payload.center_lat)
        center_lng, center_lat = round(center_lng, 6), round(center_lat, 6)
    search_keywords = payload.keywords or ([payload.keyword] if payload.keyword else [])
    center_name = payload.center_name or ("当前位置" if has_coordinates else "")
    center_address = payload.center_address or payload.center_name or ("浏览器定位" if has_coordinates else "")
    keyword_label = "、".join(search_keywords)[:100]
    radius_m = int(payload.radius_km * 1000)
    if is_around_search:
        cached = find_cached_task(
            db,
            center_name=center_name,
            center_address=center_address,
            city=payload.city,
            keyword=keyword_label,
            keywords=search_keywords,
            types_code=payload.types_code,
            radius_m=radius_m,
            providers=effective_providers,
            ttl_hours=settings.cache_ttl_hours,
            center_lng=center_lng,
            center_lat=center_lat,
        )
        if cached is not None:
            return _task_read(cached, cache_hit=True)

    unavailable = [name for name in payload.providers if not settings.key_for(name)]
    if unavailable:
        if unavailable == ["amap"]:
            detail = (
                "高德地图 Web 服务 Key 尚未配置。"
                "请在服务器 .env 中填写 AMAP_KEY，保存后重启服务。"
            )
        else:
            detail = f"以下地图平台尚未配置服务器 Key：{', '.join(unavailable)}"
        raise HTTPException(
            status_code=400,
            detail=detail,
        )

    task = Task(
        center_name=center_name or None,
        center_address=center_address or None,
        keyword=keyword_label if is_around_search else None,
        types_code=payload.types_code or None,
        radius_m=radius_m if is_around_search else None,
        center_lng=center_lng,
        center_lat=center_lat,
        city=payload.city,
        keywords_json=json.dumps(
            search_keywords,
            ensure_ascii=False,
        ),
        providers_json=json.dumps(effective_providers, ensure_ascii=False),
    )
    db.add(task)
    db.commit()
    db.refresh(task)
    background_tasks.add_task(collect_task, task.id, payload.pages)
    return _task_read(task)


@router.get("", response_model=list[TaskRead])
def list_tasks(limit: int = Query(20, ge=1, le=100), db: Session = Depends(get_db)):
    tasks = db.scalars(select(Task).order_by(Task.id.desc()).limit(limit)).all()
    return [_task_read(task) for task in tasks]


@pois_router.get("/pois", response_model=list[PoiRead])
def list_pois(
    keyword: str = Query("", max_length=100, description="匹配机构名称或分类"),
    city: str = Query("", max_length=100),
    provider: str = Query("", pattern="^(amap|baidu|tencent)?$"),
    product: str = Query("", max_length=50),
    verified: bool | None = Query(None),
    limit: int = Query(100, ge=1, le=500),
    db: Session = Depends(get_db),
):
    query = select(Poi)
    keyword = keyword.strip()
    city = city.strip()
    if keyword:
        pattern = f"%{keyword}%"
        query = query.where(or_(Poi.name.ilike(pattern), Poi.category.ilike(pattern)))
    if city:
        query = query.where(Poi.city == city)
    if provider:
        query = query.where(Poi.provider == provider)
    if product.strip():
        query = query.where(Poi.products_json.ilike(f'%"{product.strip()}"%'))
    if verified is not None:
        query = query.where(Poi.products_verified == verified)
    query = query.order_by(
        case((Poi.distance_m.is_(None), 1), else_=0),
        Poi.distance_m,
        Poi.id,
    ).limit(limit)
    return [_poi_read(row) for row in db.scalars(query).all()]


@pois_router.patch("/pois/{poi_id}/products", response_model=PoiRead)
def update_poi_products(
    poi_id: int,
    payload: PoiProductUpdate,
    db: Session = Depends(get_db),
):
    poi = db.get(Poi, poi_id)
    if poi is None:
        raise HTTPException(status_code=404, detail="机构不存在")
    poi.products_json = json.dumps(payload.products, ensure_ascii=False)
    poi.products_verified = payload.verified
    poi.product_note = payload.note or None
    poi.product_updated_at = datetime.utcnow()
    db.commit()
    db.refresh(poi)
    return _poi_read(poi)


@router.get("/{task_id}", response_model=TaskRead)
def get_task(task_id: int, db: Session = Depends(get_db)):
    task = db.get(Task, task_id)
    if task is None:
        raise HTTPException(status_code=404, detail="任务不存在")
    return _task_read(task)


@router.get("/{task_id}/results", response_model=PoiPage)
def get_results(
    task_id: int,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    db: Session = Depends(get_db),
):
    if db.get(Task, task_id) is None:
        raise HTTPException(status_code=404, detail="任务不存在")
    condition = Poi.task_id == task_id
    total = db.scalar(select(func.count(Poi.id)).where(condition)) or 0
    rows = db.scalars(
        select(Poi)
        .where(condition)
        .order_by(
            case((Poi.distance_m.is_(None), 1), else_=0),
            Poi.distance_m,
            Poi.id,
        )
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    return PoiPage(
        items=[_poi_read(row) for row in rows],
        total=total,
        page=page,
        page_size=page_size,
    )


@router.get("/{task_id}/pois", response_model=list[PoiRead])
def get_task_pois(
    task_id: int,
    limit: int = Query(500, ge=1, le=1000),
    db: Session = Depends(get_db),
):
    if db.get(Task, task_id) is None:
        raise HTTPException(status_code=404, detail="任务不存在")
    rows = db.scalars(
        select(Poi)
        .where(Poi.task_id == task_id)
        .order_by(
            case((Poi.distance_m.is_(None), 1), else_=0),
            Poi.distance_m,
            Poi.id,
        )
        .limit(limit)
    ).all()
    return [_poi_read(row) for row in rows]


@router.get("/{task_id}/export")
def export_results(task_id: int, db: Session = Depends(get_db)):
    task = db.get(Task, task_id)
    if task is None:
        raise HTTPException(status_code=404, detail="任务不存在")
    rows = db.scalars(
        select(Poi)
        .where(Poi.task_id == task_id)
        .order_by(
            case((Poi.distance_m.is_(None), 1), else_=0),
            Poi.distance_m,
            Poi.id,
        )
    ).all()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "采集结果"
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "周边机构采集结果"
    sheet["A1"].font = Font(name="Arial", size=14, bold=True, color="000000")
    sheet["A2"] = "任务 ID"
    sheet["B2"] = task.id
    sheet["D2"] = "任务状态"
    sheet["E2"] = task.status
    sheet["A3"] = "目的地"
    sheet["B3"] = task.center_name or ""
    sheet["D3"] = "半径（米）"
    sheet["E3"] = task.radius_m
    sheet["A4"] = "目的地地址"
    sheet["B4"] = task.center_address or ""
    sheet["D4"] = "中心坐标"
    sheet["E4"] = (
        f"{task.center_lng}, {task.center_lat}"
        if task.center_lng is not None and task.center_lat is not None
        else ""
    )
    sheet["A5"] = "关键词"
    sheet["B5"] = task.keyword or "、".join(json.loads(task.keywords_json))
    sheet["D5"] = "结果数量"
    sheet["E5"] = len(rows)

    for row in range(2, 6):
        sheet[f"A{row}"].font = Font(name="Arial", bold=True, color="526575")
        sheet[f"D{row}"].font = Font(name="Arial", bold=True, color="526575")

    header_row = 7
    headers = [
        "序号", "来源", "关键词", "名称", "分类", "地址", "省", "市", "区县",
        "联系电话", "联系邮箱", "官方网站", "距离（米）", "经营产品",
        "产品已核实", "核实备注", "经度", "纬度", "采集时间",
    ]
    for column, value in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=column, value=value)
    for index, poi in enumerate(rows, start=1):
        sheet.append([
            index, poi.provider, poi.keyword, poi.name, poi.category, poi.address,
            poi.province, poi.city, poi.district, poi.phone, poi.email, poi.website,
            poi.distance_m, "、".join(poi.products),
            "是" if poi.products_verified else "否", poi.product_note,
            poi.longitude, poi.latitude, poi.created_at,
        ])

    header_fill = PatternFill("solid", fgColor="17324D")
    white_side = Side(style="thin", color="FFFFFF")
    for cell in sheet[header_row]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=white_side, right=white_side)
    sheet.row_dimensions[header_row].height = 24

    last_row = header_row + len(rows)
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=last_row):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="17324D")
            cell.alignment = Alignment(vertical="center")
        row[4].alignment = Alignment(vertical="top", wrap_text=True)
        row[5].alignment = Alignment(vertical="top", wrap_text=True)
        row[18].number_format = "yyyy-mm-dd hh:mm:ss"

    widths = [8, 11, 16, 28, 25, 42, 12, 12, 14, 18, 28, 32, 13, 26, 12, 28, 13, 13, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=header_row, column=index).column_letter].width = width
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:S{last_row}"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    place = task.center_name or task.city or "采集结果"
    filename = quote(f"POI任务_{task_id}_{place}.xlsx")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )
